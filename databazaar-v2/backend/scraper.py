import time
import re
import os
import hashlib
import threading
from io import BytesIO
import pandas as pd
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options

SCROLL_TIMES = 10
WAIT_TIME = 2

# ── In-memory live frame buffer for WebSocket streaming ──
# Stores the latest JPEG frame per job_id: { job_id: { "data": base64_str, "hash": md5_hex } }
LIVE_FRAMES: dict[int, dict] = {}
_frame_lock = threading.Lock()

def push_scraper_frame(driver, job_id):
    """Capture a compressed JPEG frame and store it in the in-memory buffer"""
    if not job_id:
        return
    try:
        # Get screenshot as PNG bytes from Selenium
        png_bytes = driver.get_screenshot_as_png()

        # Compress to JPEG using Pillow for smaller frames (~20-50KB vs ~500KB PNG)
        from PIL import Image as PILImage
        img = PILImage.open(BytesIO(png_bytes))

        # Resize to max 800px width to keep frames lightweight
        max_width = 800
        if img.width > max_width:
            ratio = max_width / img.width
            img = img.resize((max_width, int(img.height * ratio)), PILImage.LANCZOS)

        buffer = BytesIO()
        img.save(buffer, format="JPEG", quality=40, optimize=True)
        jpeg_bytes = buffer.getvalue()

        import base64
        b64_data = base64.b64encode(jpeg_bytes).decode("utf-8")
        frame_hash = hashlib.md5(jpeg_bytes).hexdigest()

        with _frame_lock:
            LIVE_FRAMES[job_id] = {"data": b64_data, "hash": frame_hash}
    except Exception:
        pass

def get_live_frame(job_id: int) -> dict | None:
    """Get the latest frame for a job (thread-safe)"""
    with _frame_lock:
        return LIVE_FRAMES.get(job_id)

def clear_scraper_frame(job_id: int):
    """Clean up frame buffer when a job finishes"""
    with _frame_lock:
        LIVE_FRAMES.pop(job_id, None)

# Backward-compatible alias — all existing call sites use this name
save_scraper_screenshot = push_scraper_frame

# ── Manual Job Interruption / Stop Registry ──
STOPPED_JOBS: set[int] = set()
_stop_lock = threading.Lock()

def stop_scraper_job(job_id: int):
    """Signal a job to stop scraping immediately"""
    with _stop_lock:
        STOPPED_JOBS.add(job_id)

def is_job_stopped(job_id: int) -> bool:
    """Check if job has received a stop signal"""
    if not job_id:
        return False
    with _stop_lock:
        return job_id in STOPPED_JOBS

def clear_job_stop(job_id: int):
    """Clean up stop registry entry"""
    with _stop_lock:
        STOPPED_JOBS.discard(job_id)



def setup_driver(headless=True):
    """Setup Chrome in headless background mode with Selenium"""
    options = Options()
    options.add_argument("--headless=new")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--disable-gpu")
    options.add_argument("--disable-blink-features=AutomationControlled")
    options.add_experimental_option("excludeSwitches", ["enable-automation"])
    options.add_experimental_option("useAutomationExtension", False)
    options.add_argument("--window-size=1400,900")
    options.add_argument(
        "user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    )
    driver = webdriver.Chrome(options=options)
    driver.execute_script("Object.defineProperty(navigator, 'webdriver', {get: () => undefined})")
    return driver

def extract_phone(text):
    """Extract Bangladesh phone number from page source or elements"""
    patterns = [
        r'(\+?88001[3-9]\d{8})',
        r'(\+?8801[3-9]\d{8})',
        r'(8801[3-9]\d{8})',
        r'(01[3-9]\d{8})',
        r'(1[3-9]\d{8})',
    ]
    for pattern in patterns:
        match = re.search(pattern, text.replace('-', '').replace(' ', ''))
        if match:
            num = match.group(1)
            if num.startswith('+88001'):
                num = '+8801' + num[6:]
            elif num.startswith('88001'):
                num = '+8801' + num[5:]
            elif num.startswith('01'):
                num = '+88' + num
            elif num.startswith('1') and len(num) == 10:
                num = '+880' + num
            elif not num.startswith('+'):
                num = '+' + num
            return num
    return ''

def scroll_results(driver, scroll_times=10, log_cb=print):
    """Scroll Google Maps sidebar panel"""
    try:
        panel = driver.find_element(By.CSS_SELECTOR, 'div[role="feed"]')
        for i in range(scroll_times):
            driver.execute_script("arguments[0].scrollTop = arguments[0].scrollHeight", panel)
            time.sleep(1.5)
            log_cb(f"Scrolling maps sidebar feed... {i+1}/{scroll_times}")
    except Exception as e:
        log_cb(f"Sidebar scroll skipped: {e}")

def scrape_query(driver, query, log_cb=print, job_id=None):
    """Scrape a search query from Google Maps"""
    results = []
    log_cb(f"Starting Google Maps search query: '{query}'")

    url = f"https://www.google.com/maps/search/{query.replace(' ', '+')}"
    driver.get(url)
    time.sleep(3)
    save_scraper_screenshot(driver, job_id)

    # Scroll sidebar
    scroll_results(driver, SCROLL_TIMES, log_cb)
    time.sleep(2)
    save_scraper_screenshot(driver, job_id)



    listings = driver.find_elements(By.CSS_SELECTOR, 'a[href*="/maps/place/"]')
    log_cb(f"Discovered {len(listings)} business elements on the page.")

    seen_names = set()

    for i, listing in enumerate(listings):
        if is_job_stopped(job_id):
            log_cb("⏹️ Interruption / Stop signal received. Halting scrape for current query...")
            break

        try:
            name = listing.get_attribute("aria-label") or ""
            href = listing.get_attribute("href") or ""

            if not name or name in seen_names:
                continue
            seen_names.add(name)

            # Click to load details
            driver.execute_script("arguments[0].click();", listing)
            time.sleep(2)
            save_scraper_screenshot(driver, job_id)

            phone = ''
            address = ''
            website = ''
            rating = ''
            category = ''

            # Phone Extraction
            try:
                phone_els = driver.find_elements(
                    By.CSS_SELECTOR,
                    'button[data-item-id*="phone"], button[aria-label*="phone"], [data-tooltip*="Copy phone"]'
                )
                for el in phone_els:
                    text = el.get_attribute("aria-label") or el.text
                    phone = extract_phone(text)
                    if phone:
                        break
            except Exception:
                pass

            # Address Extraction
            try:
                addr_els = driver.find_elements(
                    By.CSS_SELECTOR, 'button[data-item-id*="address"], [data-tooltip*="Copy address"]'
                )
                for el in addr_els:
                    txt = el.get_attribute("aria-label") or el.text
                    if txt and len(txt) > 5:
                        address = txt.replace("Address: ", "").strip()
                        break
            except Exception:
                pass

            # Website Extraction
            try:
                web_els = driver.find_elements(By.CSS_SELECTOR, 'a[data-item-id*="authority"]')
                for el in web_els:
                    w = el.get_attribute("href") or ""
                    if w and "google" not in w:
                        website = w
                        break
            except Exception:
                pass

            # Rating
            try:
                rating_el = driver.find_element(By.CSS_SELECTOR, 'span[aria-label*="stars"]')
                rating = rating_el.get_attribute("aria-label") or ""
            except Exception:
                pass

            # Category
            try:
                cat_el = driver.find_element(By.CSS_SELECTOR, 'button[jsaction*="category"]')
                category = cat_el.text.strip()
            except Exception:
                pass

            entry = {
                "Name": name,
                "Phone": phone,
                "Address": address,
                "Website": website,
                "Rating": rating,
                "Category": category,
                "Maps URL": href,
                "Query": query,
            }
            results.append(entry)
            status_tag = "[OK]" if phone else "[NO PHONE]"
            log_cb(f"Scraped listing: {name[:30]} {status_tag} {phone}")
            if i % 2 == 0:
                save_scraper_screenshot(driver, job_id)



        except Exception as e:
            log_cb(f"Error extracting listing index {i+1}: {e}")
            continue

    return results

def save_to_excel(all_results, filename):
    """Save results to format-aligned Excel spreadsheet"""
    if not all_results:
        return
    df = pd.DataFrame(all_results)
    if "Phone" in df.columns:
        df["Phone"] = df["Phone"].astype(str).str.replace(r'\.0$', '', regex=True).replace({'nan': '', 'None': ''})
    df = df.drop_duplicates(subset=["Name", "Phone"])

    # Sort with phone numbers first
    df["has_phone"] = df["Phone"].apply(lambda x: 0 if x else 1)
    df = df.sort_values(["has_phone", "Name"]).drop(columns=["has_phone"])

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Coaching Centers")
        ws = writer.sheets["Coaching Centers"]

        col_widths = {
            "A": 35, "B": 18, "C": 40, "D": 30, "E": 10, "F": 20, "G": 50, "H": 35
        }
        for col, width in col_widths.items():
            ws.column_dimensions[col].width = width

        from openpyxl.styles import PatternFill, Font, Alignment
        header_fill = PatternFill(start_color="0D1B3E", end_color="0D1B3E", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        green_fill = PatternFill(start_color="E1F5EE", end_color="E1F5EE", fill_type="solid")
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            if row[1].value:  # has phone
                for cell in row:
                    cell.fill = green_fill
